from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(__file__).resolve().parents[1] / "public" / "Asim_Aslah_GenAI_AIResearch_FullStack_Resume.docx"
FONT = "Arial"
INK = RGBColor(0x12, 0x18, 0x20)
MUTED = RGBColor(0x4B, 0x55, 0x63)
LINK = "1D4ED8"


def set_run_font(run, size, *, bold=False, color=INK):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def set_paragraph(paragraph, *, before=0, after=0, line=1.05, keep=False):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.keep_with_next = keep


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), LINK)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.append(color)
    run_properties.append(underline)
    run.append(run_properties)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def section_heading(document, text):
    paragraph = document.add_paragraph()
    set_paragraph(paragraph, before=7, after=2.5, line=1, keep=True)
    run = paragraph.add_run(text.upper())
    set_run_font(run, 10.4, bold=True)
    return paragraph


def entry_heading(document, title, meta):
    paragraph = document.add_paragraph()
    set_paragraph(paragraph, before=3, after=1, line=1.02, keep=True)
    run = paragraph.add_run(title)
    set_run_font(run, 10.15, bold=True)
    run = paragraph.add_run(f" | {meta}")
    set_run_font(run, 9.65, bold=True, color=MUTED)
    return paragraph


def body_paragraph(document, text):
    paragraph = document.add_paragraph()
    set_paragraph(paragraph, after=2, line=1.08)
    run = paragraph.add_run(text)
    set_run_font(run, 9.4)
    return paragraph


def skill_line(document, label, text):
    paragraph = document.add_paragraph()
    set_paragraph(paragraph, after=1.3, line=1.04)
    label_run = paragraph.add_run(f"{label}: ")
    set_run_font(label_run, 9.25, bold=True)
    text_run = paragraph.add_run(text)
    set_run_font(text_run, 9.25)
    return paragraph


def bullet(document, text):
    paragraph = document.add_paragraph(style="List Bullet")
    set_paragraph(paragraph, after=1.2, line=1.05)
    paragraph.paragraph_format.left_indent = Inches(0.18)
    paragraph.paragraph_format.first_line_indent = Inches(-0.14)
    run = paragraph.add_run(text)
    set_run_font(run, 9.15)
    return paragraph


def configure_styles(document):
    normal = document.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(9.65)
    normal.font.color.rgb = INK

    title = document.styles["Title"]
    title.font.name = FONT
    title._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    title._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    title.font.size = Pt(19)
    title.font.bold = True
    title.font.color.rgb = INK
    title_properties = title._element.get_or_add_pPr()
    title_border = title_properties.find(qn("w:pBdr"))
    if title_border is not None:
        title_properties.remove(title_border)


def build_resume():
    document = Document()
    configure_styles(document)
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.4)
    section.bottom_margin = Inches(0.4)
    section.left_margin = Inches(0.62)
    section.right_margin = Inches(0.62)
    section.header_distance = Inches(0.2)
    section.footer_distance = Inches(0.2)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph(title, after=1.5, line=1)
    title.add_run("ASIM ASLAH P M")

    role = document.add_paragraph()
    role.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph(role, after=2, line=1)
    set_run_font(role.add_run("AI/ML & Generative AI Developer | AI Research | Full-Stack Software Engineer"), 10.5, bold=True)

    contact = document.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph(contact, after=5, line=1.05)
    set_run_font(contact.add_run("+91 9207900426 | asimaslu7@gmail.com | Palakkad, Kerala, India\n"), 8.9, color=MUTED)
    add_hyperlink(contact, "LinkedIn", "https://www.linkedin.com/in/asim-aslah-pm-05906a222")
    set_run_font(contact.add_run(" | "), 8.9, color=MUTED)
    add_hyperlink(contact, "GitHub", "https://github.com/AsimAslah")
    set_run_font(contact.add_run(" | "), 8.9, color=MUTED)
    add_hyperlink(contact, "Portfolio", "https://asim-portfolio.asimaslu7.workers.dev/")

    section_heading(document, "Professional Summary")
    body_paragraph(
        document,
        "M.Tech candidate in Artificial Intelligence and Data Science with industry experience building AI-enabled web applications. Develops Python and FastAPI services, integrates pretrained PyTorch models, and ships tested React applications. Current research explores conformal privacy-risk routing for retrieval-augmented generation under PII distribution shift.",
    )

    section_heading(document, "Technical Skills")
    skill_line(document, "Languages", "Python, TypeScript, JavaScript, SQL")
    skill_line(document, "Generative AI and Research", "RAG research, privacy-risk estimation, conformal risk control research, distribution shift, NLP, named entity recognition")
    skill_line(document, "Machine Learning", "PyTorch, computer vision, object detection, semantic segmentation, TripoSR, YOLOv8, BiSeNet, OpenCV, NumPy, Pandas")
    skill_line(document, "Backend and Data", "FastAPI, REST APIs, Uvicorn, Supabase, Cloudflare D1")
    skill_line(document, "Frontend and Cloud", "React, Next.js, HTML, CSS, Liquid, Shopify, Cloudflare Workers")
    skill_line(document, "Testing and Tools", "pytest, Node.js test runner, GitHub Actions, Git, GitHub, ESLint, Vite")

    section_heading(document, "Professional Experience")
    entry_heading(document, "Software Development Intern, Realviz LLP (Brickrat)", "Malappuram, Kerala | Jun 2026 - Jul 2026")
    bullet(document, "Built the application workflow around pretrained TripoSR, converting one furniture image into downloadable OBJ and GLB assets.")
    bullet(document, "Developed FastAPI endpoints for upload, preprocessing, inference, validation, preview, and download, extending an initial Gradio prototype into a complete web application.")
    bullet(document, "Integrated Supabase to persist generated-model metadata and public asset URLs for storage and retrieval.")
    bullet(document, "Implemented browser 3D previews, PWA behavior, Apple Quick Look USDZ delivery, and Android Scene Viewer paths.")

    section_heading(document, "Selected Projects")
    entry_heading(document, "DataVeil", "Python, FastAPI, Presidio, spaCy, YOLOv8, OpenCV, BiSeNet, pytest")
    bullet(document, "Built a local-first application that masks sensitive text, documents, and images before external generative AI use, then restores protected content from user-held recovery data.")
    bullet(document, "Combined Presidio, spaCy, YOLOv8, OpenCV, BiSeNet, eight masking effects, ORB/ECC restoration, bounded privacy controls, and a 62-test GitHub Actions suite.")

    entry_heading(document, "ASIMON Fashion Studio", "Shopify, Liquid, JavaScript, HTML, CSS")
    bullet(document, "Developed a completed Shopify storefront with merchant-editable sections, responsive catalog layouts, variant-ready product pages, native cart behavior, and accessible motion controls.")

    entry_heading(document, "Asim Portfolio", "Next.js, React, TypeScript, Cloudflare Workers, D1")
    bullet(document, "Built and deployed this tested Next.js portfolio with case studies, BYTE guide, accessible controls, Cloudflare D1 feedback, and SEO metadata.")

    section_heading(document, "Research Experience")
    entry_heading(document, "PIIShift-RAG", "Ongoing M.Tech Research")
    bullet(document, "Investigating privacy-risk routing for retrieval-augmented generation when PII patterns shift between development and deployment.")
    bullet(document, "The proposed method combines PII detection, risk estimation, conformal calibration, and policy-aware routing; implementation and empirical evaluation remain ongoing.")

    section_heading(document, "Education")
    entry_heading(document, "Master of Technology in Artificial Intelligence and Data Science", "Government Engineering College Palakkad | 2025 - 2027")
    body_paragraph(document, "APJ Abdul Kalam Technological University | Machine learning, computer vision, pattern recognition, statistics, and optimization")
    entry_heading(document, "Bachelor of Technology in Computer Science and Engineering", "MES College of Engineering | 2020 - 2024")
    body_paragraph(document, "APJ Abdul Kalam Technological University")

    properties = document.core_properties
    properties.title = "Asim Aslah GenAI AI Research Full Stack Resume"
    properties.subject = "Professional resume"
    properties.author = "Asim Aslah P M"
    properties.keywords = "Python, Generative AI, RAG, AI Research, FastAPI, PyTorch, React, TypeScript, Cloudflare"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_resume()
